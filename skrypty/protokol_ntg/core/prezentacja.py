# -*- coding: utf-8 -*-
"""NTG - prezentacja PPTX ze wzoru ze znacznikami (bez Qt-owych okien).

Wzór (plik poza wtyczką, wskazywany w oknie) to zwykła prezentacja:
slajdy stałe (podstawy prawne, procedura, zdjęcia...) zostają bez zmian,
a w polach tekstowych stoją znaczniki {{nazwa}}:
- tekst (w dowolnym miejscu pola, także w tabeli wzoru): {{miejscowosc}},
  {{data_narady}}, {{tytul}}, {{nr_umowy}}, {{data_umowy}}, {{rok_od}},
  {{rok_do}}, {{starostwo}}, {{gminy}} (lista nazw), {{gmina}} (na slajdzie
  gminy),
- tabela (pole zawierające TYLKO znacznik; tabela w miejscu i szerokości
  pola, wielkość czcionki ze znacznika, za długa - kolejne slajdy z
  "(cd.)" w tytule): {{tab_admin}}, {{tab_tsl}}, {{tab_td}}, {{tab_reb}},
  {{tab_hod}}, {{tab_wiek}}, {{tab_niezal}}, {{tab_ochr}}, {{tab_stat}},
  {{tab_roz}}, {{tab_etaty}} - dla gminy; z końcówką _calosc (np.
  {{tab_etaty_calosc}}) - cały obiekt, wszystkie gminy w jednej tabeli,
- mapa (pole z samym znacznikiem; obraz wpasowany w pole z zachowaniem
  proporcji): {{mapa_ogolna}} oraz dla gminy {{mapa_tsl}},
  {{mapa_zabiegi}}, {{mapa_fop}} - PNG z "Generuj mapy".
Slajd ze znacznikiem gminy (tabela bez _calosc, mapa gminy, {{gmina}})
jest powtarzany dla każdej gminy (kolejność kodów); bez {{gmina}} na
slajdzie nazwa gminy jest dopisywana do tytułu. Gmina bez danych w żadnej
tabeli slajdu (i bez map na nim) jest pomijana. Pola bez znaczników (np.
miejsca na zdjęcia) - bez zmian.

Biblioteka python-pptx (MIT) dołączona w ../biblioteki (QGIS jej nie ma;
lxml i Pillow są w QGIS).
"""
import copy
import math
import os
import re
import sys

_BIBL = os.path.join(os.path.dirname(os.path.dirname(
    os.path.abspath(__file__))), "biblioteki")
if _BIBL not in sys.path:
    sys.path.insert(0, _BIBL)

from pptx import Presentation  # noqa: E402
from pptx.enum.text import PP_ALIGN  # noqa: E402
from pptx.opc.constants import RELATIONSHIP_TYPE as RT  # noqa: E402
from pptx.util import Emu, Pt  # noqa: E402

ZNACZNIK = re.compile(r"\{\{\s*([A-Za-z_]+)\s*\}\}")
NS_R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
LICZBA = re.compile(r"^-?\d[\d ]*(,\d+)?$|^-$")
CZCIONKA = 12          # pt - gdy znacznik nie ma ustawionej wielkości
SZER_ZNAKU = 0.55      # średnia szerokość znaku / wielkość czcionki
MARGINES = Pt(2)       # margines komórki (góra/dół, lewo/prawo x2)

# tabela -> kolumny [(nagłówek, klucz)]; gm - kolumna gminy (bez niej na
# slajdzie gminy), suma - kolumna sumowana w wierszu Razem gminy (wiersze
# Razem całego obiektu wtedy odpadają), bez_razem - same wiersze gmin,
# pary - układ w kilku parach kolumn obok siebie
TABELE = {
    "tab_admin": {"kol": [
        ("Nazwa gminy", "naz"), ("Nr obr. ewid.", "obr"),
        ("Nazwa obr. ewid.", "obr_naz"), ("Pow. użytków Ls w ha", "pow_ls"),
        ("Pow. użytków do Lz (>0,1 ha) w ha", "pow_lz"),
        ("Pow. lasów pozaewidencyjnych w ha", "pow_nie"),
        ("Powierzchnia do rozliczenia w ha", "pow_roz")], "gm": "naz"},
    "tab_tsl": {"kol": [
        ("Nazwa gminy", "gm"), ("TSL", "tsl"),
        ("Powierzchnia w ha", "pow")], "gm": "gm", "suma": "pow"},
    "tab_td": {"kol": [
        ("Nazwa gminy", "gm"), ("Typ siedliskowy lasu (TSL)", "tsl"),
        ("Typ drzewostanu (TD)", "td"), ("Powierzchnia w ha", "pow")],
        "gm": "gm", "suma": "pow"},
    "tab_reb": {"kol": [
        ("Nazwa gminy", "gm"), ("Rodzaj zabiegu", "zab"),
        ("Sumaryczna powierzchnia zabiegu w ha", "pow")],
        "gm": "gm", "suma": "pow"},
    "tab_hod": {"kol": [
        ("Nazwa gminy", "gm"), ("Rodzaj zabiegu", "zab"),
        ("Sumaryczna powierzchnia zabiegu w ha", "pow")],
        "gm": "gm", "suma": "pow"},
    "tab_wiek": {"kol": [("Gatunek", "gat"), ("Wiek rębności", "wiek")],
                 "pary": 4},
    "tab_niezal": {"kol": [
        ("Nazwa gminy", "gm"), ("Rodzaj powierzchni", "rodz"),
        ("Nazwa", "naz"), ("Powierzchnia w ha", "pow")],
        "gm": "gm", "suma": "pow"},
    "tab_ochr": {"kol": [
        ("Gmina", "gm"), ("Oznaczenie FOP", "ozn"),
        ("Nazwa i/lub kod FOP", "naz"),
        ("Powierzchnia graficzna FOP w ha", "powg"),
        ("Sumaryczna powierzchnia całych wydzieleń w obrębie FOP w ha",
         "wydz")], "gm": "gm"},
    "tab_stat": {"kol": [
        ("Nazwa gminy", "gm"), ("Ilość wydz.", "lwydz"),
        ("Średnia pow. wydz.", "pwydz"), ("Ilość jedn. rej. G", "grej"),
        ("Ilość działek", "ldz"), ("Średnia pow. dz.", "pdz"),
        ("Średnia pow. Ls", "pls"),
        ("Średni ważony przyrost bieżący", "pwaz")],
        "gm": "gm", "bez_razem": True},
    "tab_roz": {"kol": [
        ("Nazwa gminy", "gm"), ("Rodzaj powierzchni", "rodz"),
        ("Powierzchnia w ha", "pow")], "gm": "gm", "suma": "pow"},
    "tab_etaty": {"kol": [
        ("Nr gminy", "gm"), ("Nazwa gminy", "naz"),
        ("ETAT_OKW ostatnia klasa wieku", "e1"),
        ("ETAT_POKW z dwóch ostatnich klas wieku", "e2"),
        ("ETAT_WSK wg potrzeb hodowlanych", "e3"),
        ("ETAT_POZ pozostałe użytki rębne nie objęte etatem", "e4"),
        ("Masa TW i TP", "e5")], "gm": ("gm", "naz"), "bez_razem": True},
}
# mapa gminy -> część nazwy PNG z "Generuj mapy"
# (<TERYT gminy>_<nazwa>_<TSL|Zabiegi|FOP>_NTG.png)
MAPY_GMIN = {"mapa_tsl": "TSL", "mapa_zabiegi": "Zabiegi", "mapa_fop": "FOP"}
MAPA_OGOLNA = "Mapa_ogolna_NTG.png"
TEKSTY = ("miejscowosc", "data_narady", "tytul", "nr_umowy", "data_umowy",
          "rok_od", "rok_do", "starostwo", "gminy", "gmina")


def _na_liczbe(tekst):
    try:
        return float(str(tekst).replace(" ", "").replace(",", "."))
    except ValueError:
        return 0.0


def _gminny(nazwa):
    """Znacznik wymagający powtórzenia slajdu dla każdej gminy."""
    return nazwa == "gmina" or nazwa in MAPY_GMIN or nazwa in TABELE


def wiersze_tabeli(nazwa, wiersze, gmina):
    """(nagłówki, [(komórki, pogrubiony)]) do slajdu; gmina=True - tabela
    jednej gminy (bez kolumny gminy, z własnym wierszem Razem)."""
    spec = TABELE[nazwa]
    kol = spec["kol"]
    gm = spec.get("gm", ())
    gm = (gm,) if isinstance(gm, str) else gm
    if gmina:
        widoczne = [k for k in kol if k[1] not in gm]
        if spec.get("suma") or spec.get("bez_razem"):
            dane = [w for w in wiersze if not w.get("razem")]
        else:
            dane = list(wiersze)
    else:
        widoczne = list(kol)
        dane = list(wiersze)
    wyn = []
    for w in dane:
        komorki = [str(w.get(k, "")) for _n, k in widoczne]
        if gmina and w.get("razem") and not komorki[0]:
            # etykieta "Razem:" była w ukrytej kolumnie gminy
            komorki[0] = next((str(w[k]) for k in gm if w.get(k)), "Razem")
        wyn.append((komorki, bool(w.get("razem"))))
    if gmina and spec.get("suma") and dane:
        i = [k for _n, k in widoczne].index(spec["suma"])
        suma = sum(_na_liczbe(w.get(spec["suma"])) for w in dane)
        komorki = ["Razem"] + [""] * (len(widoczne) - 1)
        komorki[i] = ("%.4f" % suma).replace(".", ",")
        wyn.append((komorki, True))
    naglowki = [n for n, _k in widoczne]
    pary = spec.get("pary")
    if pary and wyn:
        n = math.ceil(len(wyn) / pary)
        uklad = []
        for r in range(n):
            komorki = []
            for p in range(pary):
                i = p * n + r
                komorki += wyn[i][0] if i < len(wyn) else [""] * len(naglowki)
            uklad.append((komorki, False))
        wyn = uklad
        naglowki = naglowki * pary
    return naglowki, wyn


# ---------------------------------------------------------------- slajdy

def _ksztalty(kontener):
    for sh in kontener:
        yield sh
        if sh.shape_type == 6:       # MSO_SHAPE_TYPE.GROUP
            yield from _ksztalty(sh.shapes)


def _ramki_tekstu(slajd):
    for sh in _ksztalty(slajd.shapes):
        if sh.has_text_frame:
            yield sh.text_frame
        if getattr(sh, "has_table", False) and sh.has_table:
            for wiersz in sh.table.rows:
                for kom in wiersz.cells:
                    yield kom.text_frame


def znaczniki(slajd):
    return {m.group(1) for tf in _ramki_tekstu(slajd)
            for m in ZNACZNIK.finditer(tf.text)}


def _znacznik_pola(sh):
    """Nazwa znacznika, gdy pole zawiera tylko jeden znacznik."""
    if not sh.has_text_frame:
        return None
    m = ZNACZNIK.fullmatch(sh.text_frame.text.strip())
    return m.group(1) if m else None


def _zamien_teksty(slajd, wartosci):
    """Znaczniki tekstu w akapitach (znacznik rozbity na kilka fragmentów
    tekstu - wynik w pierwszym, formatowanie pierwszego)."""
    for tf in _ramki_tekstu(slajd):
        for p in tf.paragraphs:
            runy = p.runs
            if not runy:
                continue
            caly = "".join(r.text for r in runy)
            if "{{" not in caly:
                continue
            nowy = ZNACZNIK.sub(
                lambda m: str(wartosci[m.group(1)])
                if m.group(1) in wartosci else m.group(0), caly)
            if nowy != caly:
                runy[0].text = nowy
                for r in runy[1:]:
                    r.text = ""


def _dopisz_do_tytulu(slajd, tekst):
    tytul = slajd.shapes.title
    if tytul is None or not tytul.text_frame.paragraphs:
        return
    p = tytul.text_frame.paragraphs[-1]
    if p.runs:
        p.runs[-1].text += tekst
    else:
        p.add_run().text = tekst


def _sld_id(prs, slajd):
    for sld_id in prs.slides._sldIdLst:
        if prs.part.related_part(sld_id.rId) is slajd.part:
            return sld_id
    return None


def duplikuj(prs, zrodlo, pozycja):
    """Kopia slajdu (kształty, tło, obrazy - te same pliki multimediów) na
    pozycji pozycja (indeks w prezentacji). Notatki i animacje nie są
    kopiowane."""
    nowy = prs.slides.add_slide(zrodlo.slide_layout)
    drzewo = nowy.shapes._spTree
    for sh in list(nowy.shapes):
        drzewo.remove(sh._element)
    mapa = {}
    for rid, rel in zrodlo.part.rels.items():
        if rel.reltype in (RT.SLIDE_LAYOUT, RT.NOTES_SLIDE):
            continue
        if rel.is_external:
            mapa[rid] = nowy.part.relate_to(rel.target_ref, rel.reltype,
                                            is_external=True)
        else:
            mapa[rid] = nowy.part.relate_to(rel.target_part, rel.reltype)

    def _przepnij(el):
        for e in el.iter():
            for atr, wart in list(e.attrib.items()):
                if atr.startswith(NS_R) and wart in mapa:
                    e.set(atr, mapa[wart])
        return el

    ext = drzewo.find("{http://schemas.openxmlformats.org/presentationml/"
                      "2006/main}extLst")
    for el in zrodlo.shapes._spTree.iterchildren():
        if el.tag.endswith("}nvGrpSpPr") or el.tag.endswith("}grpSpPr") \
                or el.tag.endswith("}extLst"):
            continue
        kopia = _przepnij(copy.deepcopy(el))
        if ext is not None:
            ext.addprevious(kopia)
        else:
            drzewo.append(kopia)
    tlo = zrodlo._element.cSld.bg
    if tlo is not None:
        nowy._element.cSld.insert(0, _przepnij(copy.deepcopy(tlo)))
    lista = prs.slides._sldIdLst
    el = lista[-1]
    lista.remove(el)
    lista.insert(pozycja, el)
    return nowy


def usun_slajd(prs, slajd):
    sld_id = _sld_id(prs, slajd)
    if sld_id is None:
        return
    rid = sld_id.rId
    prs.slides._sldIdLst.remove(sld_id)
    prs.part.drop_rel(rid)


# ---------------------------------------------------------------- tabela

def _rozmiar_czcionki(sh):
    for p in sh.text_frame.paragraphs:
        for r in p.runs:
            if r.font.size:
                return r.font.size.pt
    return CZCIONKA


def _szerokosc_tekstu(tekst, pt, pogrub=False):
    """EMU; pogrubiony tekst jest ok. 10% szerszy."""
    return len(tekst) * pt * SZER_ZNAKU * 12700 * (1.1 if pogrub else 1)


def _szerokosci_kolumn(naglowki, wiersze, szer, pt):
    """Szerokość kolumn proporcjonalnie do treści (bez zawijania, jeśli się
    da), nagłówek - co najmniej najdłuższe słowo."""
    minimum, wagi = [], []      # bez zawijania danych / całego nagłówka
    for i, nag in enumerate(naglowki):
        tresc = max([_szerokosc_tekstu(w[0][i], pt, w[1]) for w in wiersze]
                    + [pt * 12700 * 2]) + 4 * MARGINES
        slowo = max(_szerokosc_tekstu(s, pt, True)
                    for s in nag.split() or [""]) + 4 * MARGINES
        minimum.append(tresc)
        wagi.append(max(tresc, slowo))
    if sum(wagi) <= szer or sum(minimum) >= szer:
        # wszystko się mieści (rozciągnięcie) albo nawet dane się nie
        # mieszczą - proporcjonalnie
        baza = wagi if sum(wagi) <= szer else minimum
        szer_k = [int(w * szer / sum(baza)) for w in baza]
    else:
        # dane bez zawijania, reszta miejsca - nagłówkom (zawijają się)
        zapas = szer - sum(minimum)
        brak = sum(w - m for w, m in zip(wagi, minimum))
        szer_k = [int(m + (w - m) * zapas / brak)
                  for w, m in zip(wagi, minimum)]
    szer_k[-1] += szer - sum(szer_k)
    return szer_k


def _wysokosc_wiersza(komorki, szer_k, pt, pogrub=False):
    linie = 1
    for tekst, sz in zip(komorki, szer_k):
        uzyteczna = max(sz - 4 * MARGINES, 1)
        linie = max(linie, math.ceil(_szerokosc_tekstu(tekst, pt, pogrub) /
                                     uzyteczna) or 1)
    return int(linie * pt * 1.2 * 12700 + 2 * MARGINES)


def strony_tabeli(naglowki, wiersze, szer, wys, pt):
    """Podział wierszy na slajdy wg szacowanej wysokości (nagłówek na
    każdym). [] - tabela pusta."""
    if not wiersze:
        return [], []
    szer_k = _szerokosci_kolumn(naglowki, wiersze, szer, pt)
    h_nag = _wysokosc_wiersza(naglowki, szer_k, pt, True)
    strony, biezaca, h = [], [], h_nag
    for w in wiersze:
        hw = _wysokosc_wiersza(w[0], szer_k, pt, w[1])
        if biezaca and h + hw > wys:
            strony.append(biezaca)
            biezaca, h = [], h_nag
        biezaca.append(w)
        h += hw
    strony.append(biezaca)
    return strony, szer_k


def _wstaw_w_miejsce(sh, nowy):
    """Nowy kształt w miejscu znacznika (ta sama kolejność warstw)."""
    sh._element.addprevious(nowy._element)
    sh._element.getparent().remove(sh._element)


def wstaw_tabele(slajd, sh, naglowki, wiersze, szer_k, pt):
    wszystkie = [(naglowki, True)] + wiersze
    wys = [_wysokosc_wiersza(k, szer_k, pt, b) for k, b in wszystkie]
    ramka = slajd.shapes.add_table(
        len(wszystkie), len(naglowki), sh.left, sh.top, sh.width,
        Emu(sum(wys)))
    tab = ramka.table
    for i, sz in enumerate(szer_k):
        tab.columns[i].width = Emu(sz)
    for r, (komorki, pogrub) in enumerate(wszystkie):
        tab.rows[r].height = Emu(wys[r])
        for c, tekst in enumerate(komorki):
            kom = tab.cell(r, c)
            kom.margin_top = kom.margin_bottom = MARGINES
            kom.margin_left = kom.margin_right = MARGINES * 2
            kom.text = tekst
            p = kom.text_frame.paragraphs[0]
            if r and LICZBA.match(tekst.strip()):
                p.alignment = PP_ALIGN.RIGHT
            elif r == 0:
                p.alignment = PP_ALIGN.CENTER
            for run in p.runs:
                run.font.size = Pt(pt)
                run.font.bold = pogrub
    _wstaw_w_miejsce(sh, ramka)


def wstaw_obraz(slajd, sh, sciezka):
    from PIL import Image
    with Image.open(sciezka) as im:
        w, h = im.size
    skala = min(sh.width / w, sh.height / h)
    sz, wy = int(w * skala), int(h * skala)
    obraz = slajd.shapes.add_picture(
        sciezka, sh.left + (sh.width - sz) // 2,
        sh.top + (sh.height - wy) // 2, sz, wy)
    _wstaw_w_miejsce(sh, obraz)


def _ustaw_tekst(sh, tekst):
    p = sh.text_frame.paragraphs[0]
    if p.runs:
        p.runs[0].text = tekst
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run().text = tekst


# ---------------------------------------------------------------- mapy

def mapa_gminy(folder, gm, rodzaj):
    """PNG mapy gminy z "Generuj mapy" (<TERYT>_<nazwa>_<rodzaj>_NTG.png;
    TERYT gminy = województwo + powiat + gmina z rodzajem) albo None."""
    if not folder or not os.path.isdir(folder):
        return None
    poczatek = "".join(gm) + "_"
    koniec = "_%s_NTG.png" % rodzaj
    return next((os.path.join(folder, f) for f in sorted(os.listdir(folder))
                 if f.startswith(poczatek) and f.endswith(koniec)), None)


# ---------------------------------------------------------------- całość

def generuj(wzor, wynik_sc, teksty, tabele, tabele_gmin, nazwy_gmin,
            folder_map):
    """Prezentacja ze wzoru. teksty - {znacznik: tekst}, tabele - tabele
    całego obiektu (DaneProtokolu.policz()), tabele_gmin - {gmina:
    tabele} (policz_gminy()), nazwy_gmin - {gmina: nazwa}, folder_map -
    PNG z "Generuj mapy". Zwraca (liczba slajdów, ostrzeżenia)."""
    prs = Presentation(wzor)
    ostrz = []
    nieznane = set()
    for slajd in list(prs.slides):
        nazwy = znaczniki(slajd)
        if not nazwy:
            continue
        nieznane |= {n for n in nazwy if not (
            n in TEKSTY or n in TABELE or n in MAPY_GMIN or
            n == "mapa_ogolna" or (n.endswith("_calosc") and
                                   n[:-len("_calosc")] in TABELE))}
        gminne = any(_gminny(n) for n in nazwy)
        # gminy zawsze w kolejności TERYT (województwo, powiat, gmina z
        # rodzajem) - ta sama na slajdach tabel i map
        wersje = sorted(tabele_gmin, key="".join) if gminne else [None]
        egzemplarze = []        # [(gmina, strona, {znacznik: wiersze})]
        for gm in wersje:
            strony_zn, ma_dane = {}, False
            for sh in slajd.shapes:
                zn = _znacznik_pola(sh)
                if zn is None:
                    continue
                if zn in MAPY_GMIN or zn == "mapa_ogolna":
                    ma_dane = True
                    continue
                calosc = zn.endswith("_calosc")
                tab = zn[:-len("_calosc")] if calosc else zn
                if tab not in TABELE:
                    continue
                zrodlo = tabele if (calosc or gm is None) else \
                    tabele_gmin[gm]
                nag, wiersze = wiersze_tabeli(
                    tab, zrodlo.get(tab, []), gmina=not calosc and
                    gm is not None)
                strony, szer_k = strony_tabeli(
                    nag, wiersze, sh.width, sh.height, _rozmiar_czcionki(sh))
                ma_dane = ma_dane or bool(strony)
                strony_zn[sh.shape_id] = (nag, strony, szer_k)
            if gm is not None and strony_zn and not ma_dane:
                ostrz.append("Slajd \"%s\": gmina %s bez danych - pominięta"
                             % (_tytul(slajd), nazwy_gmin.get(gm, gm)))
                continue
            n = max([len(s[1]) for s in strony_zn.values()] + [1])
            egzemplarze += [(gm, i, strony_zn) for i in range(n)]
        if not egzemplarze:
            usun_slajd(prs, slajd)
            continue
        if len(egzemplarze) == 1 and egzemplarze[0][1] == 0:
            _wypelnij(slajd, *egzemplarze[0], nazwy, teksty, nazwy_gmin,
                      folder_map, ostrz)
            continue
        poz = list(prs.slides._sldIdLst).index(_sld_id(prs, slajd))
        for j, egz in enumerate(egzemplarze, 1):
            kopia = duplikuj(prs, slajd, poz + j)
            _wypelnij(kopia, *egz, nazwy, teksty, nazwy_gmin, folder_map,
                      ostrz)
        usun_slajd(prs, slajd)
    if nieznane:
        ostrz.append("Nieznane znaczniki we wzorze (zostawione bez zmian): "
                     + ", ".join("{{%s}}" % n for n in sorted(nieznane)))
    # add_slide nadaje nazwę slideN.xml wg liczby slajdów - po kopiach i
    # usunięciach nazwy się powtarzają (zdublowane pliki w zip-ie)
    prs.part.rename_slide_parts([s.rId for s in prs.slides._sldIdLst])
    prs.save(wynik_sc)
    return len(prs.slides), ostrz


def _tytul(slajd):
    t = slajd.shapes.title
    return t.text_frame.text.strip() if t is not None else "bez tytułu"


def _wypelnij(slajd, gm, strona, strony_zn, nazwy, teksty, nazwy_gmin,
              folder_map, ostrz):
    """Egzemplarz slajdu: tabele (strona podziału), mapy, teksty."""
    # identyfikatory kształtów kopii slajdu są te same co we wzorze
    for sh in list(slajd.shapes):
        zn = _znacznik_pola(sh)
        if zn is None:
            continue
        if sh.shape_id in strony_zn:
            nag, strony, szer_k = strony_zn[sh.shape_id]
            if not strony:
                _ustaw_tekst(sh, "Brak danych" if strona == 0 else "")
                if strona:
                    sh._element.getparent().remove(sh._element)
            elif strona < len(strony):
                wstaw_tabele(slajd, sh, nag, strony[strona], szer_k,
                             _rozmiar_czcionki(sh))
            else:
                sh._element.getparent().remove(sh._element)
            continue
        if zn == "mapa_ogolna" or zn in MAPY_GMIN:
            if strona:
                sh._element.getparent().remove(sh._element)
                continue
            if zn == "mapa_ogolna":
                sc = os.path.join(folder_map or "", MAPA_OGOLNA)
                sc = sc if os.path.isfile(sc) else None
                oczek = MAPA_OGOLNA
            else:
                sc = mapa_gminy(folder_map, gm, MAPY_GMIN[zn])
                oczek = "%s_..._%s_NTG.png" % ("".join(gm), MAPY_GMIN[zn])
            if sc:
                wstaw_obraz(slajd, sh, sc)
            else:
                _ustaw_tekst(sh, "Brak mapy: " + oczek)
                ostrz.append("Brak mapy %s w %s" % (oczek, folder_map))
    wartosci = dict(teksty)
    if gm is not None:
        wartosci["gmina"] = nazwy_gmin.get(gm, "".join(gm))
    _zamien_teksty(slajd, wartosci)
    if gm is not None and "gmina" not in nazwy:
        _dopisz_do_tytulu(slajd, " - " + wartosci["gmina"])
    if strona:
        _dopisz_do_tytulu(slajd, " (cd.)")
