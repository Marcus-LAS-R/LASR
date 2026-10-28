# -*- coding: utf-8 -*-
"""PPOŻ - kompleksy leśne pow. 300 ha (bez Qt-owych okien).

Kompleks wyznaczany bezwzględnie (bez względu na własność): nasze WYDZ +
wydzielenia BDL (PGL LP i poza PGL LP wg PUL), wszystkie rodzaje
powierzchni. Poligony oddalone o <= odległość łączenia (także łańcuchem
przez inne) należą do jednego kompleksu (union-find na indeksie
przestrzennym, odległość liczona dokładnie, nie buforem). Powierzchnia
kompleksu = pole sumy samych poligonów (bez przerw między nimi); zostają
kompleksy > PROG_HA, w których są nasze wydzielenia.

Dane z BDL dociągane rundami: start = zasięg WYDZ + zapas; kompleks z naszymi
wydzieleniami, który ma jeszcze <= PROG_HA i nie mieści się (z odległością
łączenia) w pobranym obszarze, dostaje kolejne komórki wokół siebie - inaczej
obcięty kompleks mógłby niesłusznie wypaść. Kompleks, który już przekroczył
próg, nie jest dociągany (duże kompleksy PGL LP mają tysiące wydzieleń) -
jeśli wychodzi poza pobrany obszar, ma NIEPELNY = 'T' (kontur i powierzchnia
obcięte do pobranego obszaru).

Powierzchnie wydzieleń w tabeli - z bazy (F_SUBAREA.SUB_AREA po
ADRESS_FOREST), nazwy obrębów - z F_COMMUNITY.
"""
import os
import re
from datetime import datetime

from qgis.core import (
    QgsCoordinateReferenceSystem, QgsCoordinateTransform, QgsFeature,
    QgsField, QgsFields, QgsGeometry, QgsProject, QgsRectangle,
    QgsSpatialIndex, QgsVectorLayer, QgsWkbTypes,
)
from qgis.PyQt.QtCore import QVariant

from ... import funkcje
from ...baza_wrapper import Baza
from . import bdl

EPSG = "EPSG:2180"
PROG_HA = 300.0
ODLEGLOSC_DOMYSLNA = 20.0
ZAPAS_DOMYSLNY = 2000.0          # m wokół WYDZ przy pierwszym pobraniu BDL
ZAPAS_DOCIAGANIA = 500.0         # m wokół kompleksu przy dociąganiu
MAX_RUND = 8
MAX_KOMOREK = 400                # 2x2 km -> 1600 km2

NASZE = "NASZE"
NAZWA_KOMPLEKSY = "KOMPLEKSY_300HA"
NAZWA_WYDZ = "WYDZ_W_KOMPLEKSACH"
NAZWA_BDL = "LASY_BDL"
FOLDER_OSOBNO = "Kompleksy_osobno"
NAZWA_XLSX = "Kompleksy_300ha.xlsx"
POLA_WYMAGANE = ("COUNTY", "DISTRICT", "MUNICIP", "COMMUNITY", "ODDZ",
                 "WYDZ", "ADR_LES")


class Przerwano(Exception):
    pass


# ---------------------------------------------------------------- WYDZ

def wczytaj_wydz(sciezka):
    """(warstwa, [(fid, geometria EPSG:2180)]) - geometrie naprawione."""
    lyr = QgsVectorLayer(sciezka, "WYDZ", "ogr")
    if not lyr.isValid():
        raise RuntimeError("Nie da się otworzyć warstwy " + sciezka)
    brak = [p for p in POLA_WYMAGANE if lyr.fields().indexOf(p) < 0]
    if brak:
        raise RuntimeError("Warstwa %s nie ma pól: %s"
                           % (os.path.basename(sciezka), ", ".join(brak)))
    crs = QgsCoordinateReferenceSystem(EPSG)
    tr = None
    if lyr.crs().isValid() and lyr.crs() != crs:
        tr = QgsCoordinateTransform(lyr.crs(), crs,
                                    QgsProject.instance().transformContext())
    geom = []
    for f in lyr.getFeatures():
        g = QgsGeometry(f.geometry())
        if g.isNull() or g.isEmpty():
            continue
        if tr is not None:
            g.transform(tr)
        g = bdl.napraw(g)
        if g is None:
            continue
        geom.append((f.id(), g))
    if not geom:
        raise RuntimeError("Warstwa %s nie ma geometrii"
                           % os.path.basename(sciezka))
    return lyr, geom


# ---------------------------------------------------------------- grupowanie

class Grupowanie:
    """Union-find poligonów, łączy pary o odległości <= odl. Poligony można
    dokładać partiami (dociąganie BDL) - wcześniejsze połączenia zostają."""

    def __init__(self, odl):
        self.odl = odl
        self.idx = QgsSpatialIndex()
        self.geom = []           # id -> QgsGeometry
        self.zrodlo = []         # id -> NASZE / PGL / INNE
        self.dane = []           # id -> fid WYDZ albo krotka BDL
        self.rodzic = []

    def znajdz(self, i):
        while self.rodzic[i] != i:
            self.rodzic[i] = self.rodzic[self.rodzic[i]]
            i = self.rodzic[i]
        return i

    def _polacz(self, a, b):
        a, b = self.znajdz(a), self.znajdz(b)
        if a != b:
            self.rodzic[max(a, b)] = min(a, b)

    def dodaj(self, poligony, postep=None):
        """poligony: [(zrodlo, geometria, dane)]."""
        nowe = []
        for zrodlo, g, dane in poligony:
            i = len(self.geom)
            self.geom.append(g)
            self.zrodlo.append(zrodlo)
            self.dane.append(dane)
            self.rodzic.append(i)
            f = QgsFeature(i)
            f.setGeometry(g)
            self.idx.addFeature(f)
            nowe.append(i)
        for k, i in enumerate(nowe):
            if postep is not None and k % 200 == 0 and \
                    not postep(k, len(nowe)):
                raise Przerwano()
            g = self.geom[i]
            ramka = QgsRectangle(g.boundingBox())
            ramka.grow(self.odl)
            for j in self.idx.intersects(ramka):
                if j == i or self.znajdz(i) == self.znajdz(j):
                    continue
                if g.distance(self.geom[j]) <= self.odl + 1e-6:
                    self._polacz(i, j)

    def grupy(self):
        """{korzeń: [id]}"""
        wyn = {}
        for i in range(len(self.geom)):
            wyn.setdefault(self.znajdz(i), []).append(i)
        return wyn

    def ramka(self, ids):
        r = QgsRectangle(self.geom[ids[0]].boundingBox())
        for i in ids[1:]:
            r.combineExtentWith(self.geom[i].boundingBox())
        return r

    def pole_ha(self, ids):
        return QgsGeometry.unaryUnion(
            [self.geom[i] for i in ids]).area() / 10000.0


def _pokrycie(komorki):
    return QgsGeometry.unaryUnion(
        [QgsGeometry.fromRect(bdl.prostokat_komorki(k)) for k in komorki])


def _miesci_sie(gr, ids, pokrycie):
    r = gr.ramka(ids)
    r.grow(gr.odl)
    return pokrycie.contains(QgsGeometry.fromRect(r))


def wyznacz(wydz_geom, odl, zapas, postep=None, log=None):
    """Zwraca (Grupowanie, lista kompleksów, liczba pobranych komórek,
    czy osiągnięto limit, niepewne). Kompleks = dict(ids, pole, niepelny);
    niepewne = [(pole, liczba naszych)] - kompleksy z naszymi wydzieleniami
    <= PROG_HA, których nie udało się dociągnąć do końca (limit)."""
    log = log or (lambda t: None)
    gr = Grupowanie(odl)
    gr.dodaj([(NASZE, g, fid) for fid, g in wydz_geom])

    zasieg = QgsRectangle(wydz_geom[0][1].boundingBox())
    for _, g in wydz_geom[1:]:
        zasieg.combineExtentWith(g.boundingBox())
    zasieg.grow(zapas)
    pobrane, widziane = set(), set()
    do_pobrania = bdl.komorki(zasieg)
    limit = False

    for runda in range(1, MAX_RUND + 1):
        if len(pobrane) + len(do_pobrania) > MAX_KOMOREK:
            do_pobrania = set(sorted(do_pobrania)[:MAX_KOMOREK - len(pobrane)])
            limit = True
        nowe = []
        zadania = [(kom, w) for kom in sorted(do_pobrania)
                   for w in sorted(bdl.WARSTWY)]
        wyniki = bdl.pobierz_rownolegle(
            zadania, lambda k, n: postep is None or postep(
                "Pobieranie z BDL (runda %d): %d/%d zapytań"
                % (runda, k, n)))
        if wyniki is None:
            raise Przerwano()
        for kom, warstwa in zadania:
            zrodlo = bdl.WARSTWY[warstwa]
            for sid, g, adres, typ, wlasc in wyniki[(kom, warstwa)]:
                if (zrodlo, sid) in widziane:
                    continue
                widziane.add((zrodlo, sid))
                nowe.append((zrodlo, g, (sid, adres, typ, wlasc)))
        pobrane |= do_pobrania
        log("Runda %d: pobrano %d komórek, nowych wydzieleń BDL: %d"
            % (runda, len(do_pobrania), len(nowe)))

        def post_gr(k, n):
            return postep is None or postep(
                "Grupowanie (runda %d): %d/%d" % (runda, k, n))
        gr.dodaj(nowe, post_gr)

        pokrycie = _pokrycie(pobrane)
        do_pobrania = set()
        for ids in gr.grupy().values():
            if not any(gr.zrodlo[i] == NASZE for i in ids):
                continue
            if _miesci_sie(gr, ids, pokrycie) or gr.pole_ha(ids) > PROG_HA:
                continue
            r = gr.ramka(ids)
            r.grow(odl + ZAPAS_DOCIAGANIA)
            do_pobrania |= bdl.komorki(r) - pobrane
        if not do_pobrania or limit:
            break
    else:
        limit = True

    pokrycie = _pokrycie(pobrane)
    kompleksy, niepewne = [], []
    for ids in gr.grupy().values():
        nasze = sum(1 for i in ids if gr.zrodlo[i] == NASZE)
        if not nasze:
            continue
        pole = gr.pole_ha(ids)
        if pole <= PROG_HA:
            if not _miesci_sie(gr, ids, pokrycie):
                niepewne.append((pole, nasze))
            continue
        kompleksy.append({"ids": ids, "pole": pole,
                          "niepelny": not _miesci_sie(gr, ids, pokrycie)})
    kompleksy.sort(key=lambda k: -k["pole"])
    for nr, k in enumerate(kompleksy, 1):
        k["nr"] = nr
    return gr, kompleksy, len(pobrane), limit, niepewne


def powierzchnie(gr, kompleks):
    """Uzupełnia kompleks o kontur i pola (ha): nasze, PGL bez naszych,
    inne (reszta)."""
    ids = kompleks["ids"]
    kontur = QgsGeometry.unaryUnion([gr.geom[i] for i in ids])
    nasze = QgsGeometry.unaryUnion(
        [gr.geom[i] for i in ids if gr.zrodlo[i] == NASZE])
    pgl = [gr.geom[i] for i in ids if gr.zrodlo[i] == "PGL"]
    pow_pgl = QgsGeometry.unaryUnion(pgl).difference(nasze).area() / 1e4 \
        if pgl else 0.0
    kompleks["kontur"] = kontur
    kompleks["pole"] = kontur.area() / 1e4
    kompleks["pow_nasze"] = nasze.area() / 1e4
    kompleks["pow_pgl"] = pow_pgl
    kompleks["pow_inne"] = max(0.0, kompleks["pole"] -
                               kompleks["pow_nasze"] - pow_pgl)


# ---------------------------------------------------------------- baza

def dane_z_bazy(baza_sc):
    """({ADRESS_FOREST: SUB_AREA}, {(COUNTY, DISTRICT, MUNICIP, COMMUNITY):
    COMMUNITY_NAME})."""
    baza = Baza(baza_sc)
    if not baza.polacz():
        raise RuntimeError("Nie udało się połączyć z bazą " + baza_sc)
    try:
        pow_w = baza.pobierz(
            "select a.ADRESS_FOREST, s.SUB_AREA from F_ARODES as a inner "
            "join F_SUBAREA as s on a.ARODES_INT_NUM = s.ARODES_INT_NUM "
            "where a.ARODES_TYP_CD = 'WYDZIEL';")
        obreby = baza.pobierz(
            "select COUNTY_CD, DISTRICT_CD, MUNICIPALITY_CD, COMMUNITY_CD, "
            "COMMUNITY_NAME from F_COMMUNITY;")
    finally:
        baza.zamknij()
    if pow_w is False or obreby is False:
        raise RuntimeError("Nie udało się odczytać F_ARODES/F_SUBAREA/"
                           "F_COMMUNITY z bazy " + baza_sc)
    powierzchnie_w = {}
    for adr, pw in pow_w:
        adr = _tekst(adr)
        if adr and adr not in powierzchnie_w:
            powierzchnie_w[adr] = pw
    nazwy = {tuple(str(x or "").strip() for x in w[:4]): (w[4] or "").strip()
             for w in obreby}
    return powierzchnie_w, nazwy


def _tekst(w):
    return "" if w is None or str(w) == "NULL" else str(w).strip()


def _klucz_naturalny(t):
    return [(0, int(c), "") if c.isdigit() else (1, 0, c)
            for c in re.findall(r"\d+|\D+", t)]


def id_obrebu(f):
    """'281409_5.0010' z pól COUNTY/DISTRICT/MUNICIP/COMMUNITY."""
    m = _tekst(f["MUNICIP"])
    return "%s%s%s_%s.%s" % (_tekst(f["COUNTY"]), _tekst(f["DISTRICT"]),
                             m[:2], m[2:], _tekst(f["COMMUNITY"]))


def wiersze_kompleksu(lyr, gr, kompleks, pow_w, nazwy):
    """Jeden wiersz na adres leśny: (obręb, nazwa obrębu, oddział,
    wydzielenie, adres, powierzchnia z bazy albo None)."""
    fidy = [gr.dane[i] for i in kompleks["ids"] if gr.zrodlo[i] == NASZE]
    wiersze = {}
    for f in lyr.getFeatures(fidy):
        adr = _tekst(f["ADR_LES"])
        if adr in wiersze:
            continue
        klucz = (_tekst(f["COUNTY"]), _tekst(f["DISTRICT"]),
                 _tekst(f["MUNICIP"]), _tekst(f["COMMUNITY"]))
        pw = pow_w.get(adr)
        wiersze[adr] = (id_obrebu(f), nazwy.get(klucz, ""),
                        _tekst(f["ODDZ"]), _tekst(f["WYDZ"]), adr,
                        float(pw) if pw is not None else None)
    return sorted(wiersze.values(), key=lambda w: (
        w[0], _klucz_naturalny(w[2]), _klucz_naturalny(w[3])))


# ---------------------------------------------------------------- zapis

def _warstwa_pamieci(nazwa, pola, typ="MultiPolygon"):
    mem = QgsVectorLayer("%s?crs=%s" % (typ, EPSG), nazwa, "memory")
    mem.dataProvider().addAttributes(pola)
    mem.updateFields()
    return mem


def _pola_kompleksow():
    pola = QgsFields()
    pola.append(QgsField("NR_KOMPL", QVariant.Int))
    for n in ("POW_HA", "POW_PGL", "POW_INNE", "POW_NASZE"):
        pola.append(QgsField(n, QVariant.Double, "double", 12, 4))
    pola.append(QgsField("NIEPELNY", QVariant.String, "string", 1))
    return pola


def _zapisz(mem, sciezka):
    ok, kom = funkcje.zapisz_shp_utf8(mem, sciezka)
    if not ok:
        raise RuntimeError("Nie zapisano %s: %s" % (sciezka, kom))
    return sciezka


def _obiekt_kompleksu(pola, k):
    f = QgsFeature(pola)
    g = QgsGeometry(k["kontur"])
    g.convertToMultiType()
    f.setGeometry(g)
    f.setAttributes([k["nr"], round(k["pole"], 4), round(k["pow_pgl"], 4),
                     round(k["pow_inne"], 4), round(k["pow_nasze"], 4),
                     "T" if k["niepelny"] else "N"])
    return f


def _obiekty_wydz(lyr, gr, k, pola, wydz_geom):
    fidy = [gr.dane[i] for i in k["ids"] if gr.zrodlo[i] == NASZE]
    wyn = []
    for f in lyr.getFeatures(fidy):
        nf = QgsFeature(pola)
        g = QgsGeometry(wydz_geom[f.id()])
        g.convertToMultiType()
        nf.setGeometry(g)
        nf.setAttributes(f.attributes() + [k["nr"]])
        wyn.append(nf)
    return wyn


def zapisz_warstwy(folder, lyr, gr, kompleksy, wydz_geom):
    """KOMPLEKSY_300HA, WYDZ_W_KOMPLEKSACH, LASY_BDL oraz osobne pliki dla
    każdego kompleksu. Zwraca listę zapisanych ścieżek."""
    wydz_geom = dict(wydz_geom)
    zapisane = []
    pola_k = _pola_kompleksow()
    mem = _warstwa_pamieci(NAZWA_KOMPLEKSY, pola_k.toList())
    mem.dataProvider().addFeatures(
        [_obiekt_kompleksu(mem.fields(), k) for k in kompleksy])
    zapisane.append(_zapisz(mem, os.path.join(folder,
                                              NAZWA_KOMPLEKSY + ".shp")))

    pola_w = QgsFields(lyr.fields())
    pola_w.append(QgsField("NR_KOMPL", QVariant.Int))
    mem = _warstwa_pamieci(NAZWA_WYDZ, pola_w.toList())
    for k in kompleksy:
        mem.dataProvider().addFeatures(
            _obiekty_wydz(lyr, gr, k, mem.fields(), wydz_geom))
    zapisane.append(_zapisz(mem, os.path.join(folder, NAZWA_WYDZ + ".shp")))

    osobno = os.path.join(folder, FOLDER_OSOBNO)
    os.makedirs(osobno, exist_ok=True)
    for k in kompleksy:
        nazwa = "KOMPLEKS_%02d" % k["nr"]
        mem = _warstwa_pamieci(nazwa, pola_k.toList())
        mem.dataProvider().addFeatures([_obiekt_kompleksu(mem.fields(), k)])
        zapisane.append(_zapisz(mem, os.path.join(osobno, nazwa + ".shp")))
        mem = _warstwa_pamieci(nazwa + "_WYDZ", pola_w.toList())
        mem.dataProvider().addFeatures(
            _obiekty_wydz(lyr, gr, k, mem.fields(), wydz_geom))
        zapisane.append(_zapisz(mem, os.path.join(osobno,
                                                  nazwa + "_WYDZ.shp")))

    # pobrane lasy z BDL - podkład do map, z numerem kompleksu (0 = poza)
    nr = {}
    for k in kompleksy:
        for i in k["ids"]:
            nr[i] = k["nr"]
    pola_b = QgsFields()
    for n, d in (("ZRODLO", 5), ("ADR_LES", 25), ("TYP_POW", 10),
                 ("WLASC", 60)):
        pola_b.append(QgsField(n, QVariant.String, "string", d))
    pola_b.append(QgsField("NR_KOMPL", QVariant.Int))
    mem = _warstwa_pamieci(NAZWA_BDL, pola_b.toList())
    obiekty = []
    for i, zrodlo in enumerate(gr.zrodlo):
        if zrodlo == NASZE:
            continue
        _sid, adres, typ, wlasc = gr.dane[i]
        f = QgsFeature(mem.fields())
        g = QgsGeometry(gr.geom[i])
        if QgsWkbTypes.geometryType(g.wkbType()) != \
                QgsWkbTypes.PolygonGeometry:
            continue
        g.convertToMultiType()
        f.setGeometry(g)
        f.setAttributes([zrodlo, adres, typ, wlasc, nr.get(i, 0)])
        obiekty.append(f)
    mem.dataProvider().addFeatures(obiekty)
    zapisane.append(_zapisz(mem, os.path.join(folder, NAZWA_BDL + ".shp")))
    return zapisane


def zapisz_xlsx(sciezka, kompleksy, tabele):
    """Arkusz 'Zestawienie' + arkusz na każdy kompleks (tabele: {nr:
    wiersze})."""
    import openpyxl
    from openpyxl.styles import Alignment, Font

    gruba = Font(bold=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Zestawienie"
    ws.append(["Nr kompleksu", "Pow. kompleksu [ha]",
               "w tym nasze wydzielenia (graf.) [ha]", "w tym PGL LP [ha]",
               "w tym inne lasy [ha]", "Liczba wydzieleń",
               "Suma pow. wydzieleń z bazy [ha]", "Kontur pełny"])
    for k in kompleksy:
        w = tabele[k["nr"]]
        ws.append([k["nr"], round(k["pole"], 4), round(k["pow_nasze"], 4),
                   round(k["pow_pgl"], 4), round(k["pow_inne"], 4), len(w),
                   round(sum(x[5] for x in w if x[5] is not None), 4),
                   "NIE (obcięty do pobranego obszaru BDL)"
                   if k["niepelny"] else "TAK"])
    naglowki = ["Obręb ewidencyjny", "Nazwa obrębu", "Oddział", "Wydzielenie",
                "Adres leśny", "Powierzchnia wydzielenia [ha]", "Uwagi"]
    for k in kompleksy:
        ws_k = wb.create_sheet("Kompleks %d" % k["nr"])
        ws_k.append(["Kompleks leśny nr %d - pow. %.4f ha"
                     % (k["nr"], k["pole"])])
        ws_k["A1"].font = gruba
        ws_k.append([])
        ws_k.append(naglowki)
        wiersze = tabele[k["nr"]]
        for w in wiersze:
            ws_k.append(list(w) + ["" if w[5] is not None
                                   else "brak wydzielenia w bazie"])
        ws_k.append(["Razem", "", "", "", "",
                     round(sum(w[5] for w in wiersze if w[5] is not None), 4)])
        for kom in ws_k[ws_k.max_row]:
            kom.font = gruba
        for kom in ws_k[3]:
            kom.font = gruba
            kom.alignment = Alignment(wrap_text=True, vertical="center")
        for kol, szer in zip("ABCDEFG", (16, 24, 10, 12, 28, 16, 26)):
            ws_k.column_dimensions[kol].width = szer
        for wiersz in ws_k.iter_rows(min_row=4, min_col=6, max_col=6):
            for kom in wiersz:
                kom.number_format = "0.0000"
    for kom in ws[1]:
        kom.font = gruba
        kom.alignment = Alignment(wrap_text=True, vertical="center")
    for kol, szer in zip("ABCDEFGH", (12, 14, 18, 14, 14, 12, 18, 30)):
        ws.column_dimensions[kol].width = szer
    wb.save(sciezka)
    return sciezka


def zapisz_raport(folder, linie):
    kat = os.path.join(folder, "Raporty")
    os.makedirs(kat, exist_ok=True)
    sc = os.path.join(kat, "raport_kompleksy_PPOZ_%s.txt"
                      % datetime.now().strftime("%Y%m%dT%H%M%S"))
    with open(sc, "w", encoding="utf-8") as f:
        f.write("\n".join(linie))
    return sc
